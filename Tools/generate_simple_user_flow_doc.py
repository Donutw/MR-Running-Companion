import os
from datetime import datetime

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn


ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DOCS_DIR = os.path.join(ROOT, "Docs")


def _font_path_candidates():
    win = os.environ.get("WINDIR", r"C:\Windows")
    return [
        os.path.join(win, "Fonts", "msyh.ttc"),      # Microsoft YaHei
        os.path.join(win, "Fonts", "msyhbd.ttc"),
        os.path.join(win, "Fonts", "simhei.ttf"),
        os.path.join(win, "Fonts", "simsun.ttc"),
        os.path.join(win, "Fonts", "arial.ttf"),
    ]


def load_font(size: int):
    for p in _font_path_candidates():
        if os.path.exists(p):
            try:
                return ImageFont.truetype(p, size=size)
            except Exception:
                continue
    return ImageFont.load_default()


def arrow(draw: ImageDraw.ImageDraw, p1, p2, color=(20, 20, 20), width=5, head=16):
    x1, y1 = p1
    x2, y2 = p2
    draw.line((x1, y1, x2, y2), fill=color, width=width)

    import math

    ang = math.atan2(y2 - y1, x2 - x1)
    left = (
        x2 - head * math.cos(ang) + head * 0.55 * math.sin(ang),
        y2 - head * math.sin(ang) - head * 0.55 * math.cos(ang),
    )
    right = (
        x2 - head * math.cos(ang) - head * 0.55 * math.sin(ang),
        y2 - head * math.sin(ang) + head * 0.55 * math.cos(ang),
    )
    draw.polygon([p2, left, right], fill=color)


def draw_round_box(draw: ImageDraw.ImageDraw, rect, title, lines, font_title, font_body):
    x1, y1, x2, y2 = rect
    draw.rounded_rectangle(rect, radius=18, outline=(30, 30, 30), width=3, fill=(250, 250, 250))
    pad = 14
    ty = y1 + pad
    draw.text((x1 + pad, ty), title, font=font_title, fill=(0, 0, 0))
    y = ty + font_title.size + 10
    for ln in lines:
        draw.text((x1 + pad, y), ln, font=font_body, fill=(20, 20, 20))
        y += font_body.size + 6


def make_simple_flow_png(path: str):
    W, H = 1600, 1050
    img = Image.new("RGB", (W, H), (255, 255, 255))
    draw = ImageDraw.Draw(img)

    font_h1 = load_font(44)
    font_title = load_font(26)
    font_body = load_font(22)
    font_small = load_font(18)

    draw.text((40, 20), "启动 + 主场景用户操作流程（简版）", font=font_h1, fill=(0, 0, 0))

    # Row y positions
    y0 = 140
    h = 150
    box_w = 300
    gap = 70

    # Boxes (left to right)
    boxes = []
    def add(i, x, title, lines, tint=None):
        rect = (x, y0, x + box_w, y0 + h)
        if tint is None:
            tint = (250, 250, 250)
        # reuse draw_round_box with current tint by drawing separately
        # We'll draw with custom fill by manual rounded rectangle + text
        draw.rounded_rectangle(rect, radius=18, outline=(30, 30, 30), width=3, fill=tint)
        pad = 14
        ty = y0 + pad
        draw.text((x + pad, ty), title, font=font_title, fill=(0, 0, 0))
        y = ty + font_title.size + 10
        for ln in lines:
            draw.text((x + pad, y), ln, font=font_body, fill=(20, 20, 20))
            y += font_body.size + 6
        boxes.append((i, x, y0, x + box_w, y0 + h))

    add(0, 60, "Init（启动场景）", ["创建全局", "PersistentRoot"], tint=(235, 245, 255))
    add(1, 60 + (box_w + gap) * 1, "BootLoader", ["加载 Menu_Scene"], tint=(235, 245, 255))
    add(2, 60 + (box_w + gap) * 2, "Menu（主菜单）", ["等待心率连接", "显示日志/放行按钮"], tint=(245, 255, 235))
    add(3, 60 + (box_w + gap) * 3, "选择 ID + TestA/TestB", ["按钮可点", "进入测试场景"], tint=(245, 245, 255))
    add(4, 60 + (box_w + gap) * 4, "Test运行", ["开始录制", "每秒记账", "触发器请求对话"], tint=(255, 250, 230))

    # Arrows main
    for i in range(4):
        b1 = boxes[i][1]
        b2 = boxes[i + 1][1]
        arrow(draw, (b1 + box_w, y0 + h / 2), (b2, y0 + h / 2))

    # End branch row
    y1 = y0 + 250
    h2 = 155
    box_w2 = 440
    # End path
    # Left box: time over
    x_end1 = 90
    draw.rounded_rectangle((x_end1, y1, x_end1 + box_w2, y1 + h2), radius=20, outline=(30, 30, 30), width=3, fill=(255, 235, 235))
    draw.text((x_end1 + 14, y1 + 14), "时间到（EndTest）", font=font_title, fill=(0, 0, 0))
    draw.text((x_end1 + 14, y1 + 14 + font_title.size + 10), "1) 停止录制/导出 CSV", font=font_body, fill=(20, 20, 20))
    draw.text((x_end1 + 14, y1 + 14 + font_title.size + 10 + font_body.size + 6), "2) 结算面板展示平均值", font=font_body, fill=(20, 20, 20))

    x_end2 = x_end1 + box_w2 + 60
    draw.rounded_rectangle((x_end2, y1, x_end2 + box_w2, y1 + h2), radius=20, outline=(30, 30, 30), width=3, fill=(250, 250, 250))
    draw.text((x_end2 + 14, y1 + 14), "返回/下一位被试", font=font_title, fill=(0, 0, 0))
    draw.text((x_end2 + 14, y1 + 14 + font_title.size + 10), "MenuUI：可再次选 ID / 再做 A/B", font=font_body, fill=(20, 20, 20))
    draw.text((x_end2 + 14, y1 + 14 + font_title.size + 10 + font_body.size + 6), "HR连接状态仍由 PersistentRoot 保持", font=font_body, fill=(20, 20, 20))

    arrow(draw, (boxes[4][3], y0 + h / 2), (x_end1, y1 + h2 / 2))

    # Heart-rate not connected branch hint
    # Draw a small decision box under Menu
    x_dec = 60 + (box_w + gap) * 2 - 40
    y_dec = y0 + 215
    draw.rounded_rectangle((x_dec, y_dec, x_dec + 340, y_dec + 120), radius=18, outline=(180, 100, 0), width=4,
                           fill=(255, 245, 225))
    draw.text((x_dec + 16, y_dec + 14), "分支：心率是否已连接？", font=font_small, fill=(140, 70, 0))
    draw.text((x_dec + 16, y_dec + 14 + font_small.size + 6), "否：按钮不可点，只显示日志", font=font_body, fill=(20, 20, 20))
    draw.text((x_dec + 16, y_dec + 14 + font_small.size + 6 + font_body.size + 6), "是：显示按钮/编号UI，允许开始", font=font_body, fill=(20, 20, 20))

    # Legend
    lx, ly = 60, 930
    draw.rounded_rectangle((lx, ly, lx + 520, ly + 95), radius=16, outline=(30, 30, 30), width=3, fill=(252, 252, 252))
    draw.text((lx + 16, ly + 14), "颜色含义", font=font_title, fill=(0, 0, 0))
    draw.text((lx + 16, ly + 14 + font_title.size + 8), "蓝色：启动/场景切换；绿色：等待连接；黄色：测试运行；红色：结束结算", font=font_small, fill=(20, 20, 20))

    img.save(path, "PNG")


def configure_cn_fonts(document: Document):
    style = document.styles["Normal"]
    style.font.name = "微软雅黑"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(11)


def generate_docx(docx_path: str, img_flow_path: str):
    doc = Document()
    configure_cn_fonts(doc)

    doc.add_heading("VR 跑步测试系统｜简版设计文档（不写代码）", level=0)
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("说明：下面的内容只做“系统拆解 + 关系梳理 + 流程说明”，不需要贴任何代码。")

    doc.add_heading("一、如果一点代码都不沾，文档应该怎么写？", level=1)
    doc.add_paragraph("建议你把它写成“游戏系统策划书”：每个模块写清楚三件事——它负责什么、它输入什么、它输出什么（数据/状态/表现）。")
    doc.add_paragraph("你要的重点不是实现细节，而是：数据怎么流、规则怎么判定、玩家怎么走、结束时怎么结算。")

    doc.add_heading("二、系统模块（用一句话解释即可）", level=1)
    for line in [
        "全局会话层：保存被试编号与 A/B 完成状态，并提供心率读取入口。",
        "心率输入层（BLE/模拟）：持续读取最新 bpm，并提供连接状态给界面。",
        "运动指标层：根据头显位移计算距离；根据距离计算平均配速；根据心率计算卡路里。",
        "记录与导出层：按固定频率把关键指标写入 CSV，并在结束时导出。",
        "触发器哨兵：根据阈值/里程碑/冷却规则，向对话系统请求台词。",
        "对话表现层：统一仲裁优先级，控制气泡打字与动画。",
        "流程总控：管理开始/结束、启停录制与结算面板显示，并在结束时切断触发器噪音。",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("三、启动与用户操作流程图（你要求的关键图）", level=1)
    if os.path.exists(img_flow_path):
        doc.add_picture(img_flow_path, width=Inches(6.8))
    doc.add_paragraph("图说明：从 Init 启动开始，先常驻心率连接器，再进入菜单让用户等待连接；连接成功后才允许进入 TestA/TestB，结束后导出 CSV 并显示结算。")

    doc.add_heading("四、主场景可以走哪些“路”？（建议写成玩家路径）", level=1)
    doc.add_paragraph("1) 未连接心率的路：只能等待、看日志；测试按钮不可点。")
    doc.add_paragraph("2) 连接成功的路：选被试 ID → 选择 TestA 或 TestB → 点击开始 → 跑步阶段实时更新 → 时间到结束并结算 → 返回菜单。")
    doc.add_paragraph("3) 触发器的路（系统内部）：运行中距离/配速/卡路里/心率/停顿/时间都会触发台词请求；对话系统根据优先级决定谁能播。")

    doc.add_heading("五、你可以写进论文的“设计步骤叙事模板”", level=1)
    for line in [
        "先明确体验目标：实时指标 + 反馈对话 + 固定时长测试 + 可导出数据。",
        "再定义系统边界：哪些负责“输入”（心率/位置），哪些负责“规则判定”（触发器），哪些负责“表现”（对话/可视化/气泡）。",
        "然后建立数据流：心率/距离/时间 → 计算器 → 触发器 → 对话管理 → 记录器/CSV → 结算面板。",
        "最后定义运行时序：开始时统一开录制与清场；结束时先停笔与导出，再弹结算，避免数据抖动。",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.save(docx_path)


def main():
    os.makedirs(DOCS_DIR, exist_ok=True)
    png_path = os.path.join(DOCS_DIR, "simple-start-user-flow.png")
    docx_path = os.path.join(DOCS_DIR, "VR_System_Simple_Guide.docx")
    make_simple_flow_png(png_path)
    generate_docx(docx_path, png_path)
    print("OK")
    print("PNG:", png_path)
    print("DOCX:", docx_path)


if __name__ == "__main__":
    main()


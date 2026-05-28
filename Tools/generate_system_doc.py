import os
from dataclasses import dataclass
from datetime import datetime

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn


ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DOCS_DIR = os.path.join(ROOT, "Docs")


def _font_path_candidates():
    win = os.environ.get("WINDIR", r"C:\Windows")
    return [
        os.path.join(win, "Fonts", "msyh.ttc"),      # Microsoft YaHei
        os.path.join(win, "Fonts", "msyhbd.ttc"),
        os.path.join(win, "Fonts", "simhei.ttf"),    # SimHei
        os.path.join(win, "Fonts", "simsun.ttc"),    # SimSun
        os.path.join(win, "Fonts", "arial.ttf"),     # fallback (no Chinese)
    ]


def load_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for p in _font_path_candidates():
        if os.path.exists(p):
            try:
                return ImageFont.truetype(p, size=size)
            except Exception:
                continue
    return ImageFont.load_default()


@dataclass
class Box:
    x: int
    y: int
    w: int
    h: int
    title: str
    lines: list[str]

    @property
    def rect(self):
        return (self.x, self.y, self.x + self.w, self.y + self.h)

    @property
    def center(self):
        return (self.x + self.w // 2, self.y + self.h // 2)

    @property
    def left_mid(self):
        return (self.x, self.y + self.h // 2)

    @property
    def right_mid(self):
        return (self.x + self.w, self.y + self.h // 2)

    @property
    def top_mid(self):
        return (self.x + self.w // 2, self.y)

    @property
    def bottom_mid(self):
        return (self.x + self.w // 2, self.y + self.h)


def draw_box(draw: ImageDraw.ImageDraw, box: Box, font_title, font_body):
    draw.rounded_rectangle(box.rect, radius=14, outline=(30, 30, 30), width=3, fill=(250, 250, 250))
    pad = 14
    title_y = box.y + pad
    draw.text((box.x + pad, title_y), box.title, font=font_title, fill=(0, 0, 0))
    y = title_y + font_title.size + 10
    for ln in box.lines:
        draw.text((box.x + pad, y), ln, font=font_body, fill=(20, 20, 20))
        y += font_body.size + 6


def arrow(draw: ImageDraw.ImageDraw, p1, p2, color=(0, 0, 0), width=4, head=14, dash=False):
    x1, y1 = p1
    x2, y2 = p2
    if dash:
        # simple dashed: draw segments
        import math

        dx, dy = x2 - x1, y2 - y1
        dist = max(1.0, (dx * dx + dy * dy) ** 0.5)
        ux, uy = dx / dist, dy / dist
        step = 12
        gap = 10
        t = 0.0
        while t < dist - head:
            a = t
            b = min(dist - head, t + step)
            ax, ay = x1 + ux * a, y1 + uy * a
            bx, by = x1 + ux * b, y1 + uy * b
            draw.line((ax, ay, bx, by), fill=color, width=width)
            t += step + gap
    else:
        draw.line((x1, y1, x2, y2), fill=color, width=width)

    # arrow head
    import math

    ang = math.atan2(y2 - y1, x2 - x1)
    left = (x2 - head * math.cos(ang) + head * 0.55 * math.sin(ang),
            y2 - head * math.sin(ang) - head * 0.55 * math.cos(ang))
    right = (x2 - head * math.cos(ang) - head * 0.55 * math.sin(ang),
             y2 - head * math.sin(ang) + head * 0.55 * math.cos(ang))
    draw.polygon([p2, left, right], fill=color)


def make_system_architecture_png(path: str):
    W, H = 1800, 1100
    img = Image.new("RGB", (W, H), (255, 255, 255))
    draw = ImageDraw.Draw(img)

    font_h1 = load_font(48)
    font_title = load_font(28)
    font_body = load_font(22)
    font_small = load_font(18)

    draw.text((40, 28), "系统关系图（模块联系 / 数据流 / 依赖）", font=font_h1, fill=(0, 0, 0))

    boxes: dict[str, Box] = {}
    boxes["scene"] = Box(60, 150, 520, 270, "场景层（Scene Flow）", [
        "BootLoader → Menu_Scene",
        "SceneSwitcher：Menu/TestA/TestB",
        "ResultPanelController：结算面板读冻结值",
    ])
    boxes["flow"] = Box(640, 150, 520, 270, "流程层（Round Controller）", [
        "TestSceneController：Start/Run/End",
        "统一启停：计算器 + CSVLogger",
        "结束时禁用 Trigger_*（切噪音）",
    ])
    boxes["global"] = Box(1220, 90, 520, 220, "全局常驻（Session）", [
        "DataManager（DontDestroyOnLoad）",
        "participantID / TestA,B 完成标记",
        "GetCurrentBpm() 门面",
    ])
    boxes["hr"] = Box(1220, 330, 520, 220, "输入设备（Heart Rate）", [
        "HeartRateBleReader：BLE/模拟心率",
        "isConnected / currentLog / LastHeartRateBpm",
        "断线重试：直连上次地址→否则扫描",
    ])
    boxes["data"] = Box(640, 450, 520, 340, "数据层（Metrics & Logging）", [
        "DistanceCalculator：XZ 距离累计",
        "PaceCalculator：平均配速（>10m）",
        "CalorieCalculator：心率公式累加",
        "CSVLogger：每秒采样→导出 CSV",
    ])
    boxes["dialog"] = Box(60, 450, 520, 340, "表现层（Dog Dialogue）", [
        "DogDialogueManager：优先级仲裁",
        "TypeBubble：渐显+打字+渐隐/常驻",
        "DogHeartRateVisualizer：按 bpm 渐变变红",
    ])
    boxes["triggers"] = Box(60, 830, 1680, 220, "触发器哨兵（Triggers → 对话）", [
        "Trigger_HeartRate(0,冷却) / Trigger_Time(1,单次) / Trigger_Distance(2,里程碑)",
        "Trigger_Calories(2,里程碑) / Trigger_Pace(3,冷却) / Trigger_Stopping(3,冷却)",
        "统一输出：dogDialogueManager.Speak(message, priority)",
    ])

    for b in boxes.values():
        draw_box(draw, b, font_title, font_body)

    # Connections
    arrow(draw, boxes["scene"].right_mid, boxes["flow"].left_mid, color=(10, 10, 10), width=5)
    arrow(draw, boxes["flow"].right_mid, boxes["global"].left_mid, color=(10, 10, 10), width=5, dash=True)  # state report
    arrow(draw, boxes["global"].bottom_mid, boxes["hr"].top_mid, color=(10, 10, 10), width=5)
    arrow(draw, boxes["flow"].bottom_mid, boxes["data"].top_mid, color=(10, 10, 10), width=5)

    # Data reads (dashed)
    arrow(draw, boxes["hr"].left_mid, (boxes["data"].x + boxes["data"].w, boxes["data"].y + 70), color=(30, 90, 200), width=4, dash=True)
    arrow(draw, boxes["data"].left_mid, boxes["dialog"].right_mid, color=(30, 90, 200), width=4, dash=True)

    # Trigger to dialogue
    arrow(draw, (boxes["triggers"].x + 320, boxes["triggers"].y), (boxes["dialog"].x + 260, boxes["dialog"].y + boxes["dialog"].h), color=(200, 70, 30), width=5)

    # Flow disables triggers
    arrow(draw, (boxes["flow"].x + 260, boxes["flow"].y + boxes["flow"].h), (boxes["triggers"].x + 900, boxes["triggers"].y), color=(120, 0, 120), width=5, dash=True)

    # Legend
    lx, ly = 1220, 590
    draw.rounded_rectangle((lx, ly, lx + 520, ly + 180), radius=14, outline=(30, 30, 30), width=3, fill=(252, 252, 252))
    draw.text((lx + 14, ly + 12), "图例", font=font_title, fill=(0, 0, 0))
    draw.line((lx + 18, ly + 60, lx + 120, ly + 60), fill=(10, 10, 10), width=5)
    draw.polygon([(lx + 120, ly + 60), (lx + 104, ly + 52), (lx + 104, ly + 68)], fill=(10, 10, 10))
    draw.text((lx + 140, ly + 48), "实线：调用/控制", font=font_small, fill=(0, 0, 0))
    # dashed sample
    for i in range(6):
        draw.line((lx + 18 + i * 18, ly + 105, lx + 28 + i * 18, ly + 105), fill=(30, 90, 200), width=4)
    draw.polygon([(lx + 120, ly + 105), (lx + 104, ly + 97), (lx + 104, ly + 113)], fill=(30, 90, 200))
    draw.text((lx + 140, ly + 93), "虚线：数据读取/状态依赖", font=font_small, fill=(0, 0, 0))

    img.save(path, "PNG")


def make_test_flow_png(path: str):
    W, H = 1800, 1150
    img = Image.new("RGB", (W, H), (255, 255, 255))
    draw = ImageDraw.Draw(img)

    font_h1 = load_font(48)
    font_title = load_font(28)
    font_body = load_font(22)

    draw.text((40, 28), "测试流程图（从进入菜单到导出 CSV）", font=font_h1, fill=(0, 0, 0))

    lane_y = [140, 330, 520, 710, 900]
    lane_h = 170
    lanes = [
        ("用户/UI", lane_y[0]),
        ("TestSceneController", lane_y[1]),
        ("计算器（距离/配速/卡路里）", lane_y[2]),
        ("触发器/对话系统", lane_y[3]),
        ("CSVLogger/结算", lane_y[4]),
    ]

    for name, y in lanes:
        draw.rounded_rectangle((60, y, W - 60, y + lane_h), radius=14, outline=(30, 30, 30), width=3, fill=(250, 250, 250))
        draw.text((80, y + 14), name, font=font_title, fill=(0, 0, 0))

    def step(x, y, w, text, fill=(235, 245, 255)):
        draw.rounded_rectangle((x, y, x + w, y + 110), radius=12, outline=(50, 50, 50), width=3, fill=fill)
        draw.text((x + 14, y + 14), text, font=font_body, fill=(0, 0, 0))
        return (x + w, y + 55)

    x0 = 260
    dx = 250

    # Lane 1 UI
    pA = step(x0, lane_y[0] + 45, 230, "进入菜单\n等待心率连接")
    pB = step(x0 + dx, lane_y[0] + 45, 230, "选择被试编号\n进入 TestA/B")
    pC = step(x0 + 2 * dx, lane_y[0] + 45, 230, "点击开始\n开始跑步")
    pD = step(x0 + 3 * dx, lane_y[0] + 45, 230, "时间到\n显示结算面板", fill=(255, 240, 240))
    pE = step(x0 + 4 * dx, lane_y[0] + 45, 230, "返回菜单")

    # Controller
    cA = step(x0 + 2 * dx, lane_y[1] + 45, 230, "StartRunning()\n切 UI/隐藏手\nStartRecording\nStartLogging", fill=(235, 255, 235))
    cB = step(x0 + 3 * dx, lane_y[1] + 45, 230, "EndTest()\n禁用 Trigger_*\nStopRecording\nStopAndExportCSV\n设置完成标记", fill=(255, 235, 255))

    # Calculators
    mA = step(x0 + 2 * dx, lane_y[2] + 45, 230, "Distance/Pace/Cal\n开始录制\n实时更新 HUD")
    mB = step(x0 + 3 * dx, lane_y[2] + 45, 230, "停止录制\n冻结最终数值", fill=(255, 240, 240))

    # Triggers/dialogue
    tA = step(x0 + 2 * dx, lane_y[3] + 45, 230, "触发器轮询\n（冷却/单次）\n→ Speak(p)", fill=(255, 250, 230))
    tB = step(x0 + 3 * dx, lane_y[3] + 45, 230, "ForceStopSpeaking\n结束后不再触发", fill=(255, 240, 240))

    # CSV/Result
    lA = step(x0 + 2 * dx, lane_y[4] + 45, 230, "每秒采样一行\n记录 Timestamp/Dist/HR/Pace/Cal")
    lB = step(x0 + 3 * dx, lane_y[4] + 45, 230, "计算平均值\n导出 CSV\naverageHR/pace 供结算读", fill=(255, 240, 240))

    # Arrows (main timeline)
    arrow(draw, pA, (x0 + dx, lane_y[0] + 100), width=5)
    arrow(draw, pB, (x0 + 2 * dx, lane_y[0] + 100), width=5)
    arrow(draw, pC, (x0 + 3 * dx, lane_y[0] + 100), width=5)
    arrow(draw, pD, (x0 + 4 * dx, lane_y[0] + 100), width=5)

    # Cross-lane calls
    arrow(draw, (x0 + 2 * dx + 115, lane_y[0] + 155), (x0 + 2 * dx + 115, lane_y[1] + 45), width=5)
    arrow(draw, (x0 + 2 * dx + 115, lane_y[1] + 155), (x0 + 2 * dx + 115, lane_y[2] + 45), width=5)
    arrow(draw, (x0 + 2 * dx + 115, lane_y[2] + 155), (x0 + 2 * dx + 115, lane_y[3] + 45), width=5)
    arrow(draw, (x0 + 2 * dx + 115, lane_y[2] + 155), (x0 + 2 * dx + 115, lane_y[4] + 45), width=5)

    arrow(draw, (x0 + 3 * dx + 115, lane_y[1] + 155), (x0 + 3 * dx + 115, lane_y[2] + 45), width=5)
    arrow(draw, (x0 + 3 * dx + 115, lane_y[1] + 155), (x0 + 3 * dx + 115, lane_y[3] + 45), width=5)
    arrow(draw, (x0 + 3 * dx + 115, lane_y[1] + 155), (x0 + 3 * dx + 115, lane_y[4] + 45), width=5)

    img.save(path, "PNG")


def add_heading(document: Document, text: str, level: int = 1):
    p = document.add_heading(text, level=level)
    return p


def add_bullets(document: Document, items: list[str]):
    for it in items:
        document.add_paragraph(it, style="List Bullet")


def configure_cn_fonts(document: Document):
    # best-effort: set default font to a Chinese-friendly font
    style = document.styles["Normal"]
    style.font.name = "微软雅黑"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(11)


def generate_docx(docx_path: str, img_arch_path: str, img_flow_path: str):
    doc = Document()
    configure_cn_fonts(doc)

    title = doc.add_heading("VR 跑步测试系统｜系统拆解与设计文档（策划视角）", level=0)
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("说明：本文档基于当前工程脚本结构整理，不修改任何现有代码与工程结构。")

    add_heading(doc, "1. 你需要在文档里写代码吗？（建议口径）", 1)
    doc.add_paragraph("一般不需要。策划文档的目标是让读者理解：系统做什么、为什么这样做、规则如何运作、数据如何流转、如何验证。")
    add_bullets(doc, [
        "建议写：模块职责、数据字典、流程图、触发规则、配置项、异常与容错、验收与测试点。",
        "仅在必要时附“接口签名/字段名”来对齐实现（例如 bpm 字段、CSV 字段），避免贴大段代码。",
    ])

    add_heading(doc, "2. 系统总体目标与体验", 1)
    add_bullets(doc, [
        "核心体验：用户在 VR 进行固定时长跑步测试（A/B 条件），实时看到距离/配速/卡路里/心率，并由虚拟狗根据优先级进行提醒与鼓励。",
        "实验输出：结束后本地导出 CSV（每秒一行），并在结算面板展示关键统计（平均心率/平均配速等）。",
        "约束：心率来自 BLE（支持模拟模式）；DataManager 常驻跨场景；触发器必须防刷（冷却/单次）。",
    ])

    add_heading(doc, "3. 系统关系图（不同系统间联系）", 1)
    if os.path.exists(img_arch_path):
        doc.add_picture(img_arch_path, width=Inches(6.7))
        doc.add_paragraph("图：模块联系 / 数据流 / 依赖（实线=调用/控制，虚线=数据读取/状态依赖）。")
    else:
        doc.add_paragraph("（图片生成失败：未找到系统关系图 PNG）")

    add_heading(doc, "4. 测试流程图（从进入菜单到导出 CSV）", 1)
    if os.path.exists(img_flow_path):
        doc.add_picture(img_flow_path, width=Inches(6.7))
        doc.add_paragraph("图：包含 UI、总控、计算器、触发器/对话、记录器/结算 的跨模块步骤。")
    else:
        doc.add_paragraph("（图片生成失败：未找到测试流程图 PNG）")

    add_heading(doc, "5. 系统模块拆解（按策划系统组织）", 1)

    add_heading(doc, "5.1 全局会话层（Session / Meta）", 2)
    add_bullets(doc, [
        "DataManager：常驻单例（DontDestroyOnLoad）。保存 participantID、TestA/TestB 完成标记，并提供 GetCurrentBpm() 作为心率门面。",
        "意义：为“当前被试/当前回合”提供唯一权威状态，保证跨场景一致性。",
    ])
    doc.add_paragraph("完善建议（不改现有代码）：如果未来要支持多轮实验/多天数据，可在文档层面规划 SessionRecord（列表）与当前 Session 指针；实现层仍可保持 DataManager 做门面。")

    add_heading(doc, "5.2 心率采集系统（BLE / 模拟）", 2)
    add_bullets(doc, [
        "HeartRateBleReader：输出 isConnected/currentLog/LastHeartRateBpm；仅在成功解包 bpm 后视为真正连接。",
        "支持 useSimulatedHeartRate：跳过 BLE，Update 生成动态心率，便于开发与串流测试。",
        "断线重试：优先直连上次地址，失败再扫描；subscribeDelaySeconds 用于降低订阅错误概率。",
    ])

    add_heading(doc, "5.3 运动指标计算（距离/配速/卡路里）", 2)
    add_bullets(doc, [
        "DistanceCalculator：只累计 XZ 平面位移，忽略 Y 抖动；StartRecording/StopRecording 控制录制。",
        "PaceCalculator：全局平均配速；距离>10m 才计算；UI 按 updateInterval 限频刷新。",
        "CalorieCalculator：按心率/年龄/体重公式计算 kcal/min 并按帧累加；负数有保底，避免显示异常。",
    ])
    doc.add_paragraph("完善建议（数值策划）：把年龄/体重/性别等参数从“场景 Inspector 固定值”抽为配置表，并与 participantID 绑定，确保实验可复现。")

    add_heading(doc, "5.4 记录与导出（CSV）", 2)
    add_bullets(doc, [
        "CSVLogger：每 1 秒采样一行：Timestamp、Distance、HeartRate、Pace、Calories。",
        "StopAndExportCSV() 时计算 averageHR/averagePace，导出到 Application.persistentDataPath。",
        "文件命名：P{participantID}_{TestA|TestB}_{yyyyMMdd_HHmmss}.csv，避免覆盖。",
    ])

    add_heading(doc, "5.5 测试流程总控（Round Controller）", 2)
    add_bullets(doc, [
        "TestSceneController：统一管理 UI 面板显隐、手模型显隐、计时、启停计算器与 CSVLogger。",
        "结束时先冻结数据与导出 CSV，再打开结算面板，保证面板读取的是最终统计值。",
        "并会禁用同物体上所有 Trigger_* 脚本，从源头切断结束后的触发噪音。",
    ])

    add_heading(doc, "5.6 虚拟狗对话与表现（优先级仲裁）", 2)
    add_bullets(doc, [
        "DogDialogueManager：优先级数字越小越高（0紧急/1节点/2成就/3闲聊）。更高优先级可打断低优先级。",
        "TypeBubble：渐显→打字→停留→渐隐；常驻模式打完字不自动消失，需要 ForceStopSpeaking() 清场。",
        "DogHeartRateVisualizer：按 bpm 在 normalColor 与 alertColor 间平滑过渡，形成直观风险提示。",
    ])

    add_heading(doc, "5.7 触发器系统（Triggers / 哨兵）", 2)
    add_bullets(doc, [
        "Trigger_HeartRate：bpm 超阈值报警（priority=0，冷却防刷）。",
        "Trigger_Time：剩余时间阈值提醒（priority=1，单次触发；通过反射读取 currentTime/isRunning）。",
        "Trigger_Distance / Trigger_Calories：里程碑成就（priority=2，while 补触发）。",
        "Trigger_Pace / Trigger_Stopping：轻提示（priority=3，冷却防刷）。",
    ])

    add_heading(doc, "6. 如何体现“我设计整个系统的想法与步骤”", 1)
    doc.add_paragraph("建议用“设计路径”来写，而不是“代码实现过程”。下面是一套可直接放进论文/答辩的叙事模板。")
    add_heading(doc, "6.1 需求 → 系统边界", 2)
    add_bullets(doc, [
        "输入：BLE 心率（或模拟）、VR 位移（头显）。",
        "输出：实时 HUD + 虚拟狗反馈 + 结算面板 + CSV 文件。",
        "边界：本工程不做云端上传；实验数据落在本地 persistentDataPath。",
    ])
    add_heading(doc, "6.2 模块化 → 单一职责", 2)
    add_bullets(doc, [
        "把“采集/计算/触发/表现/记录/流程”拆开，避免互相硬耦合。",
        "用 DataManager 作为会话门面，解决跨场景状态与心率入口问题。",
        "触发器只做规则判定，不直接改 UI；所有台词统一走 DogDialogueManager 仲裁。",
    ])
    add_heading(doc, "6.3 运行时序 → 先冻结再展示", 2)
    add_bullets(doc, [
        "开始：总控统一开启录制与记录，确保数据源一致。",
        "进行：计算器更新指标，触发器依据指标请求对话，管理器按优先级仲裁。",
        "结束：先停笔（StopRecording/StopAndExportCSV 计算平均值），再弹结算面板读取最终值。",
    ])

    add_heading(doc, "7. 配置表建议（可选，作为未来扩展点）", 1)
    doc.add_paragraph("目前阈值/里程碑多在 Inspector。若后续要更像“系统策划”，可规划以下配置表（不要求现在改代码）：")
    doc.add_paragraph("JSON（触发器阈值/里程碑）示例：")
    doc.add_paragraph(
        '{\n'
        '  "heartRate": { "warningThreshold": 150, "cooldownSeconds": 15, "priority": 0 },\n'
        '  "timeReminder": { "thresholdSeconds": 60, "priority": 1 },\n'
        '  "distanceMilestonesMeters": [100, 500, 1000],\n'
        '  "calorieMilestonesKcal": [10, 20, 50],\n'
        '  "pace": { "slowThresholdMinPerKm": 8.0, "cooldownSeconds": 20, "priority": 3 }\n'
        '}'
    )

    doc.add_page_break()
    add_heading(doc, "附录 A：实现对齐（字段名/文件名，便于追溯）", 1)
    add_bullets(doc, [
        "入口：BootLoader.cs（加载 Menu_Scene）",
        "菜单：MenuUI.cs、SceneSwitcher.cs",
        "会话：DataManager.cs（常驻 + participantID + 完成标记）",
        "心率：HeartRateBleReader.cs（BLE/模拟 + isConnected + LastHeartRateBpm）",
        "计算器：DistanceCalculator.cs、PaceCalculator.cs、CalorieCalculator.cs",
        "记录：CSVLogger.cs（averageHR / averagePace / 导出 CSV）",
        "对话：DogDialogueManager.cs、TypeBubble.cs、DogHeartRateVisualizer.cs",
        "触发器：Trigger_*.cs",
    ])

    doc.save(docx_path)


def main():
    os.makedirs(DOCS_DIR, exist_ok=True)
    img_arch = os.path.join(DOCS_DIR, "system-architecture.png")
    img_flow = os.path.join(DOCS_DIR, "test-flow.png")
    docx_path = os.path.join(DOCS_DIR, "VR跑步测试系统_策划视角_系统拆解.docx")

    make_system_architecture_png(img_arch)
    make_test_flow_png(img_flow)
    generate_docx(docx_path, img_arch, img_flow)

    print("OK")
    print("DOCX:", docx_path)
    print("IMG1:", img_arch)
    print("IMG2:", img_flow)


if __name__ == "__main__":
    main()

